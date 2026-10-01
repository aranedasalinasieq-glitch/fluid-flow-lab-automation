from __future__ import annotations
from openpyxl import load_workbook
from pathlib import Path

def _rows(ws, header_row=4):
    headers = [c.value for c in ws[header_row]]
    out = []
    for row in ws.iter_rows(min_row=header_row+1, values_only=True):
        if all(v is None for v in row):
            continue
        out.append({headers[i]: row[i] for i in range(min(len(headers), len(row)))})
    return out

def load_input(path):
    wb = load_workbook(path, data_only=True)
    required = ["Parametros","Calibracion","Canerias","Accesorios","Pitot"]
    missing = [s for s in required if s not in wb.sheetnames]
    if missing:
        raise ValueError("Faltan hojas obligatorias: " + ", ".join(missing))
    params = {}
    for r in _rows(wb["Parametros"]):
        if r.get("Parámetro") is not None:
            params[str(r["Parámetro"]).strip()] = r.get("Valor")
    return {
        "params": params,
        "calibracion": _rows(wb["Calibracion"]),
        "canerias": _rows(wb["Canerias"]),
        "accesorios": _rows(wb["Accesorios"]),
        "pitot": _rows(wb["Pitot"]),
    }
