from __future__ import annotations
from pathlib import Path
import math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import ScatterChart, Reference, Series
import matplotlib.pyplot as plt
import numpy as np

BLUE="17365D"; MID="4472C4"; LIGHT="D9EAF7"; YELLOW="FFF2CC"

def _write_table(ws, rows):
    if not rows:
        ws.append(["Sin datos"])
        return
    keys=list(rows[0].keys())
    ws.append(keys)
    for r in rows:
        ws.append([r.get(k) for k in keys])
    for c in ws[1]:
        c.fill=PatternFill("solid", fgColor=MID); c.font=Font(color="FFFFFF",bold=True)
        c.alignment=Alignment(horizontal="center",wrap_text=True)
    ws.freeze_panes="A2"
    ws.auto_filter.ref=ws.dimensions
    for col in ws.columns:
        letter=col[0].column_letter
        width=min(24,max(10,max(len(str(x.value)) if x.value is not None else 0 for x in col)+2))
        ws.column_dimensions[letter].width=width

def create_plots(result, outdir):
    outdir=Path(outdir); outdir.mkdir(parents=True,exist_ok=True)
    cal=result["calibration"]; rows=result["calibration_rows"]
    dp=np.array([r["DeltaP_Pa"] for r in rows]); q=np.array([r["Q_m3s"] for r in rows])
    x=np.linspace(max(0,dp.min()),dp.max(),250)
    plt.figure(figsize=(7,5))
    plt.scatter(dp,q,label="Datos experimentales")
    plt.plot(x,[cal.q(v) for v in x],label=f"Q = {cal.a:.3e}√ΔP + {cal.b:.3e}\nR² = {cal.r2:.4f}")
    plt.xlabel("ΔP [Pa]"); plt.ylabel("Q [m³/s]"); plt.title("Calibración de placa orificio")
    plt.grid(True,alpha=.25); plt.legend(); plt.tight_layout()
    p1=outdir/"calibracion_Q_DeltaP.png"; plt.savefig(p1,dpi=160); plt.close()

    p2=outdir/"perfil_pitot.png"
    plt.figure(figsize=(7,5))
    angles=sorted(set(r["Apertura válvula"] for r in result["pitot"]))
    for a in angles:
        rr=sorted([r for r in result["pitot"] if r["Apertura válvula"]==a],key=lambda z:z["Posicion_m"])
        plt.plot([r["Posicion_m"] for r in rr],[r["v_local_ms"] for r in rr],marker="o",label=f"{a:g}°")
    plt.xlabel("Posición [m]"); plt.ylabel("Velocidad local [m/s]"); plt.title("Perfil de velocidades — Pitot")
    plt.grid(True,alpha=.25); plt.legend(); plt.tight_layout(); plt.savefig(p2,dpi=160); plt.close()
    return [p1,p2]

def write_results(result, path):
    wb=Workbook(); ws=wb.active; ws.title="Resumen"
    cal=result["calibration"]
    summary=[
        ["Fluid Flow Lab — Resultados automáticos","Valor"],
        ["Modelo calibración","Q = a·√ΔP + b"],
        ["a",cal.a],["b",cal.b],["R²",cal.r2],
        ["ΔP mínimo calibrado [Pa]",cal.dp_min],["ΔP máximo calibrado [Pa]",cal.dp_max],
        ["ρ aire [kg/m³]",result["rho_air"]],["Velocidad sonido [m/s]",result["speed_sound"]],
    ]
    for r in summary: ws.append(r)
    for c in ws[1]:
        c.fill=PatternFill("solid",fgColor=BLUE); c.font=Font(color="FFFFFF",bold=True)
    ws.column_dimensions["A"].width=34; ws.column_dimensions["B"].width=28
    ws2=wb.create_sheet("Advertencias")
    ws2.append(["Advertencia"]); 
    for w in result["warnings"] or ["Sin advertencias automáticas."]: ws2.append([w])
    ws2["A1"].fill=PatternFill("solid",fgColor=MID); ws2["A1"].font=Font(color="FFFFFF",bold=True)
    ws2.column_dimensions["A"].width=90
    for name,key in [("Calibracion_calc","calibration_rows"),("Canerias_calc","pipes"),
                     ("Accesorios_calc","accessories"),("Pitot_calc","pitot"),
                     ("Pitot_resumen","pitot_summary")]:
        sh=wb.create_sheet(name); _write_table(sh,result[key])
    wb.save(path)
